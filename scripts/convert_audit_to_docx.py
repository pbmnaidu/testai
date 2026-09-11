from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "MPLADS_CODEBASE_AUDIT.md"
OUTPUT = ROOT / "output" / "MPLADS_CODEBASE_AUDIT.docx"
BACKTICK = chr(96)
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\x60.+?\x60|\*.+?\*)")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Liberation Sans", size=10, bold=False, italic=False, color="1F2937") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=8, color="64748B")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_inline(paragraph, text: str, size=10, color="1F2937") -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith(BACKTICK):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Liberation Mono", size=max(size - 0.4, 7), color="334155")
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def set_paragraph_spacing(paragraph, before=0, after=5, line=1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc: Document, text: str, style="Body Text") -> None:
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text, size=9.4 if style == "Body Text" else 9.0)
    set_paragraph_spacing(paragraph, after=5, line=1.08)


def parse_table_row(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [part.strip() for part in row.split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.autofit = True
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            set_cell_shading(cell, "163A5F" if row_index == 0 else ("F3F7FB" if row_index % 2 == 0 else "FFFFFF"))
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(paragraph, after=0, line=1.0)
            value = row[col_index] if col_index < len(row) else ""
            add_inline(paragraph, value, size=7.2, color="FFFFFF" if row_index == 0 else "1F2937")
            for run in paragraph.runs:
                run.bold = row_index == 0
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(19)
    section.right_margin = Mm(19)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string("1F2937")

    title = styles["Title"]
    title.font.name = "Liberation Sans"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(7)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Liberation Sans"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string("475569")
    subtitle.paragraph_format.space_after = Pt(10)

    for name, size, before, after in [
        ("Heading 1", 15, 12, 6),
        ("Heading 2", 11.5, 9, 4),
        ("Heading 3", 10, 7, 3),
        ("Heading 4", 9.4, 5, 2),
    ]:
        style = styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Liberation Mono"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Mono")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Mono")
    code_style.font.size = Pt(7.2)
    code_style.font.color.rgb = RGBColor.from_string("334155")
    code_style.paragraph_format.left_indent = Mm(4)
    code_style.paragraph_format.right_indent = Mm(4)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("MPLADS Codebase Audit")
    set_run_font(header_run, size=8, color="64748B")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("MPLADS Codebase Audit  |  Page ")
    set_run_font(footer_run, size=8, color="64748B")
    add_page_field(footer)

    doc.core_properties.title = "MPLADS AI Monitoring Platform Codebase Audit"
    doc.core_properties.subject = "Complete technical audit and evaluator review"
    doc.core_properties.author = "Codex"


def render_markdown_to_docx() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    index = 0
    seen_title = False
    seen_subtitle = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("~~~"):
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("~~~"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            paragraph = doc.add_paragraph(style="Code Block")
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run("\n".join(code_lines))
            set_run_font(run, name="Liberation Mono", size=7.2, color="334155")
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F1F5F9")
            p_pr.append(shd)
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not seen_title:
                paragraph = doc.add_paragraph(style="Title")
                add_inline(paragraph, text, size=24, color="000000")
                seen_title = True
            elif level == 2 and seen_title and not seen_subtitle:
                paragraph = doc.add_paragraph(style="Subtitle")
                add_inline(paragraph, text, size=12, color="475569")
                seen_subtitle = True
            else:
                style_level = min(max(level - 1, 1), 4)
                paragraph = doc.add_paragraph(style=f"Heading {style_level}")
                sizes = {1: 15, 2: 11.5, 3: 10, 4: 9.4}
                add_inline(paragraph, text, size=sizes[style_level], color="000000")
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [parse_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue

        blockquote = re.match(r"^>\s?(.*)$", stripped)
        if blockquote:
            parts = [blockquote.group(1)]
            index += 1
            while index < len(lines):
                match = re.match(r"^>\s?(.*)$", lines[index].strip())
                if not match:
                    break
                parts.append(match.group(1))
                index += 1
            paragraph = doc.add_paragraph(style="Body Text")
            paragraph.paragraph_format.left_indent = Mm(5)
            paragraph.paragraph_format.right_indent = Mm(3)
            add_inline(paragraph, " ".join(parts), size=9.4, color="334155")
            for run in paragraph.runs:
                run.italic = True
            set_paragraph_spacing(paragraph, after=6, line=1.08)
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", raw)
        number = re.match(r"^\s*\d+[.)]\s+(.*)$", raw)
        if bullet or number:
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, (bullet or number).group(1), size=9.2)
            set_paragraph_spacing(paragraph, after=2, line=1.05)
            index += 1
            continue

        if re.fullmatch(r"\s*[-*_]{3,}\s*", raw):
            index += 1
            continue

        parts = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index]
            next_stripped = next_line.strip()
            if (
                not next_stripped
                or next_stripped.startswith("~~~")
                or re.match(r"^(#{1,6})\s+", next_stripped)
                or next_stripped.startswith("|")
                or re.match(r"^>\s?", next_stripped)
                or re.match(r"^\s*[-*+]\s+", next_line)
                or re.match(r"^\s*\d+[.)]\s+", next_line)
            ):
                break
            parts.append(next_stripped)
            index += 1
        add_body(doc, " ".join(parts))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    render_markdown_to_docx()
